"""Prepare a read-only, reviewable product/stock join from store exports.

Usage: python scripts/prepare-inventory-import.py CATALOG.xlsb INVENTORY.xlsx OUTPUT.json
Requires pyxlsb and openpyxl. Does not update the ERP or Feishu.
"""
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pyxlsb import open_workbook

PHONE_CATEGORIES = {"智能手机", "功能机", "旧手机"}


def text(value):
    return str(value).strip() if value is not None else ""


def number(value, label):
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label}不是数字：{value!r}")
    if result < 0:
        raise ValueError(f"{label}不能为负数：{value!r}")
    return result


def catalog_rows(path):
    with open_workbook(str(path)) as book:
        with book.get_sheet(1) as sheet:
            rows = [[cell.v for cell in row] for row in sheet.rows()]
    if not rows:
        raise ValueError("商品列表为空")
    headers = [text(value) for value in rows[0]]
    required = {"商品编号", "商品全名", "一级分类", "序列号管理", "预设售价1", "零售价"}
    if not required.issubset(headers):
        raise ValueError("商品列表缺少字段：" + ", ".join(sorted(required - set(headers))))
    result = {}
    for row_number, values in enumerate(rows[1:], 2):
        row = dict(zip(headers, values))
        code = text(row.get("商品编号"))
        if not code:
            continue
        if code in result:
            raise ValueError(f"商品编号重复：{code}（第{row_number}行）")
        row["_sourceRow"] = row_number
        result[code] = row
    return result


def inventory_rows(path):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = None
    result = []
    for row_number, values in enumerate(rows, 1):
        labels = [text(value) for value in values]
        if headers is None:
            if "商品编号" in labels and "库存数量" in labels:
                headers = labels
            continue
        row = dict(zip(headers, values))
        code = text(row.get("商品编号"))
        if not code:
            continue
        quantity = number(row.get("库存数量"), f"库存第{row_number}行数量")
        if quantity <= 0:
            continue
        if not quantity.is_integer():
            raise ValueError(f"库存第{row_number}行数量不是整数：{quantity}")
        row["_sourceRow"] = row_number
        result.append(row)
    if headers is None:
        raise ValueError("库存表缺少商品编号或库存数量表头")
    return result


def prepare(catalog_path, inventory_path):
    catalog = catalog_rows(catalog_path)
    stock = inventory_rows(inventory_path)
    seen = set()
    records = []
    for entry in stock:
        code = text(entry["商品编号"])
        if code in seen:
            raise ValueError("库存表商品编号重复：" + code)
        seen.add(code)
        product = catalog.get(code)
        if product is None:
            raise ValueError("库存商品未在商品列表找到：" + code)
        prices = [
            ("库存预设售价", entry.get("预设售价")),
            ("商品预设售价1", product.get("预设售价1")),
            ("商品零售价", product.get("零售价")),
        ]
        positive = [(label, number(value, f"{code} {label}")) for label, value in prices if value is not None and text(value)]
        chosen = next(((label, value) for label, value in positive if value > 0), ("待确认", None))
        major = text(product.get("一级分类"))
        serial_managed = major in PHONE_CATEGORIES or text(product.get("序列号管理")) == "启用"
        records.append({
            "externalCode": code,
            "name": text(product.get("商品全名")),
            "shortName": text(product.get("简名")) or text(product.get("商品全名")),
            "sourceCategory1": major,
            "sourceCategory2": text(product.get("二级分类")),
            "sourceCategory3": text(product.get("三级分类")),
            "sourceCategory4": text(product.get("四级分类")),
            "spec": text(product.get("规格")) or text(product.get("型号")),
            "unit": text(entry.get("基本单位")) or text(product.get("基本单位")) or "件",
            "quantity": int(number(entry.get("库存数量"), f"{code} 数量")),
            "costPrice": number(entry.get("成本均价"), f"{code} 成本均价"),
            "stockValue": number(entry.get("库存总价"), f"{code} 库存总价"),
            "salePrice": chosen[1],
            "priceSource": chosen[0],
            "serialManaged": serial_managed,
            "needsPriceReview": chosen[1] is None,
            "needsSerials": serial_managed,
            "catalogRow": product["_sourceRow"],
            "inventoryRow": entry["_sourceRow"],
        })
    summary = {
        "catalogProducts": len(catalog),
        "inStockProducts": len(records),
        "stockQuantity": sum(row["quantity"] for row in records),
        "stockValue": round(sum(row["stockValue"] for row in records), 2),
        "serialManagedProducts": sum(row["serialManaged"] for row in records),
        "missingSalePriceProducts": sum(row["needsPriceReview"] for row in records),
        "categoryCounts": dict(Counter(row["sourceCategory1"] for row in records)),
    }
    return {"source": {"catalog": catalog_path.name, "inventory": inventory_path.name, "status": "historical snapshot; do not post to live inventory"}, "summary": summary, "products": records}


def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    catalog_path, inventory_path, output_path = map(Path, sys.argv[1:])
    payload = prepare(catalog_path, inventory_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload["summary"], ensure_ascii=False))


if __name__ == "__main__":
    main()